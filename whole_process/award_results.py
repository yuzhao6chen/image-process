"""只读加载中标结果 Excel，并按项目解析唯一中标记录。"""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


REQUIRED_HEADERS = ("项目名称", "标段", "中标人", "金额")
MATCHED = "MATCHED"
REVIEW_REQUIRED = "REVIEW_REQUIRED"
_PROJECT_SEPARATOR_PATTERN = re.compile(r"[\s\-_\u2013\u2014]+")
_WHITESPACE_PATTERN = re.compile(r"\s+")
_NUMERIC_TEXT_PATTERN = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)$")
_NO_WINNER_VALUES = {"无数据"}


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def _normalize_header(value: Any) -> str:
    return _WHITESPACE_PATTERN.sub("", _cell_text(value))


def normalize_project_name(value: str) -> str:
    """仅忽略文件夹中常见的空白和分隔符，不做语义模糊匹配。"""
    return _PROJECT_SEPARATOR_PATTERN.sub("", _cell_text(value)).casefold()


def normalize_company_name(value: str) -> str:
    """企业身份只忽略空白和全半角差异，保留公司名称实质字符。"""
    return _WHITESPACE_PATTERN.sub("", _cell_text(value)).casefold()


def format_award_amount(value: Decimal | None) -> str | None:
    if value is None:
        return None
    text = format(value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _parse_amount(value: Any) -> tuple[Decimal | None, str | None]:
    if value is None or _cell_text(value) == "":
        return None, "金额为空"
    if isinstance(value, bool):
        return None, "金额不能是布尔值"
    if isinstance(value, float) and not math.isfinite(value):
        return None, "金额不是有限数值"

    raw = str(value) if isinstance(value, (int, float, Decimal)) else _cell_text(value)
    normalized = raw.replace(",", "").strip()
    if not isinstance(value, (int, float, Decimal)) and not _NUMERIC_TEXT_PATTERN.fullmatch(normalized):
        return None, f"金额不是纯数值：{raw}"
    try:
        amount = Decimal(normalized)
    except InvalidOperation:
        return None, f"金额无法解析：{raw}"
    if not amount.is_finite():
        return None, "金额不是有限数值"
    if amount < 0:
        return None, f"金额为负数：{raw}"
    return amount, None


@dataclass(frozen=True)
class AwardRecord:
    project_name: str
    lot_name: str | None
    winner_name: str | None
    amount: Decimal | None
    sheet_name: str
    row_number: int
    issue: str | None = None


@dataclass(frozen=True)
class AwardResolution:
    status: str
    record: AwardRecord | None
    message: str | None = None


@dataclass
class AwardRegistry:
    source_name: str
    records_by_project: dict[str, tuple[AwardRecord, ...]]
    load_error: str | None = None

    @classmethod
    def unavailable(cls, source_name: str, message: str) -> "AwardRegistry":
        return cls(source_name=source_name, records_by_project={}, load_error=message)

    def resolve(self, project_key: str) -> AwardResolution:
        if self.load_error:
            return AwardResolution(REVIEW_REQUIRED, None, self.load_error)

        normalized_project = normalize_project_name(project_key)
        if not normalized_project:
            return AwardResolution(REVIEW_REQUIRED, None, "项目目录名为空，无法匹配中标结果")

        matched_keys: list[str]
        if normalized_project in self.records_by_project:
            matched_keys = [normalized_project]
        else:
            contained = [
                key
                for key in self.records_by_project
                if len(key) >= 6 and key in normalized_project
            ]
            if not contained:
                return AwardResolution(
                    REVIEW_REQUIRED,
                    None,
                    f"中标结果表中未匹配到项目：{project_key}",
                )
            longest = max(len(key) for key in contained)
            matched_keys = [key for key in contained if len(key) == longest]
            if len(matched_keys) != 1:
                names = sorted({
                    record.project_name
                    for key in matched_keys
                    for record in self.records_by_project[key]
                })
                return AwardResolution(
                    REVIEW_REQUIRED,
                    None,
                    "项目名称匹配不唯一：" + "、".join(names),
                )

        records = self.records_by_project[matched_keys[0]]
        distinct = _deduplicate_records(records)
        if len(distinct) > 1:
            lot_matches = [
                record
                for record in distinct
                if record.lot_name
                and normalize_project_name(record.lot_name) in normalized_project
            ]
            lot_matches = _deduplicate_records(lot_matches)
            if len(lot_matches) == 1:
                distinct = lot_matches
            else:
                equivalent_outcomes = {
                    (
                        normalize_company_name(record.winner_name or ""),
                        record.amount,
                        record.issue,
                    )
                    for record in distinct
                }
                if len(equivalent_outcomes) == 1:
                    distinct = [distinct[0]]

        if len(distinct) != 1:
            locations = "、".join(
                f"{record.sheet_name}!第{record.row_number}行" for record in distinct
            )
            return AwardResolution(
                REVIEW_REQUIRED,
                None,
                f"项目对应多条不同的中标记录，无法确定唯一中标公司：{locations}",
            )

        record = distinct[0]
        if not record.winner_name:
            return AwardResolution(
                REVIEW_REQUIRED,
                None,
                record.issue
                or f"中标记录未提供中标人：{record.sheet_name}!第{record.row_number}行",
            )
        return AwardResolution(MATCHED, record, record.issue)


def _record_signature(record: AwardRecord) -> tuple[str, str, Decimal | None, str | None]:
    return (
        normalize_project_name(record.lot_name or ""),
        normalize_company_name(record.winner_name or ""),
        record.amount,
        record.issue,
    )


def _deduplicate_records(records: list[AwardRecord] | tuple[AwardRecord, ...]) -> list[AwardRecord]:
    unique: dict[tuple[str, str, Decimal | None, str | None], AwardRecord] = {}
    for record in records:
        unique.setdefault(_record_signature(record), record)
    return list(unique.values())


def _find_header_row(worksheet: Any) -> tuple[int, dict[str, int]] | None:
    max_row = min(int(worksheet.max_row or 0), 20)
    max_column = min(int(worksheet.max_column or 0), 100)
    if max_row < 1 or max_column < 1:
        return None
    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=max_row,
            min_col=1,
            max_col=max_column,
            values_only=True,
        ),
        start=1,
    ):
        positions: dict[str, int] = {}
        for index, value in enumerate(row):
            header = _normalize_header(value)
            if header in REQUIRED_HEADERS and header not in positions:
                positions[header] = index
        if all(header in positions for header in REQUIRED_HEADERS):
            return row_number, positions
    return None


def _records_from_worksheet(worksheet: Any) -> list[AwardRecord]:
    header = _find_header_row(worksheet)
    if header is None:
        return []
    header_row, positions = header
    records: list[AwardRecord] = []
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        project_name = _cell_text(row[positions["项目名称"]])
        lot_name = _cell_text(row[positions["标段"]]) or None
        winner_text = _cell_text(row[positions["中标人"]])
        amount_value = row[positions["金额"]]
        if not project_name:
            continue
        if not lot_name and not winner_text and amount_value is None:
            # 允许工作表中存在合并单元格的分组标题行。
            continue

        winner_name: str | None = winner_text or None
        issue_parts: list[str] = []
        if winner_text in _NO_WINNER_VALUES:
            winner_name = None
            issue_parts.append(
                f"中标人为“{winner_text}”：{worksheet.title}!第{row_number}行"
            )
        elif not winner_text:
            issue_parts.append(f"中标人为空：{worksheet.title}!第{row_number}行")

        amount, amount_issue = _parse_amount(amount_value)
        if amount_issue:
            issue_parts.append(f"{amount_issue}：{worksheet.title}!第{row_number}行")
        records.append(AwardRecord(
            project_name=project_name,
            lot_name=lot_name,
            winner_name=winner_name,
            amount=amount,
            sheet_name=str(worksheet.title),
            row_number=row_number,
            issue="；".join(issue_parts) or None,
        ))
    return records


def load_award_registry(path: Path) -> AwardRegistry:
    """总是返回可查询的索引；Excel 问题由企业画像标记为待核验，不阻断招标链路。"""
    source_name = path.name or "中标结果表"
    try:
        if path.is_symlink():
            return AwardRegistry.unavailable(source_name, f"中标结果表不能是符号链接：{source_name}")
        if not path.is_file():
            return AwardRegistry.unavailable(source_name, f"中标结果表不存在或不是文件：{source_name}")
    except OSError as exc:
        return AwardRegistry.unavailable(
            source_name,
            f"中标结果表无法访问：{exc.__class__.__name__}",
        )

    try:
        from openpyxl import load_workbook
    except Exception as exc:
        return AwardRegistry.unavailable(
            source_name,
            f"openpyxl 不可用，无法读取中标结果 Excel：{exc.__class__.__name__}",
        )

    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        records: list[AwardRecord] = []
        matched_sheets = 0
        for worksheet in workbook.worksheets:
            if _find_header_row(worksheet) is None:
                continue
            matched_sheets += 1
            records.extend(_records_from_worksheet(worksheet))
        if matched_sheets == 0:
            return AwardRegistry.unavailable(
                source_name,
                "中标结果 Excel 中未找到包含“项目名称、标段、中标人、金额”的表头",
            )
        if not records:
            return AwardRegistry.unavailable(source_name, "中标结果 Excel 中没有可用的项目记录")

        grouped: dict[str, list[AwardRecord]] = {}
        for record in records:
            key = normalize_project_name(record.project_name)
            if key:
                grouped.setdefault(key, []).append(record)
        return AwardRegistry(
            source_name=source_name,
            records_by_project={key: tuple(value) for key, value in grouped.items()},
        )
    except Exception as exc:
        return AwardRegistry.unavailable(
            source_name,
            f"中标结果 Excel 读取失败：{exc.__class__.__name__}",
        )
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                # 读取阶段已结束，关闭失败不应反向阻断招标链路。
                pass


__all__ = [
    "AwardRecord",
    "AwardRegistry",
    "AwardResolution",
    "MATCHED",
    "REVIEW_REQUIRED",
    "format_award_amount",
    "load_award_registry",
    "normalize_company_name",
    "normalize_project_name",
]
